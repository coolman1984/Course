#!/usr/bin/env python3
"""Builds the finished results of the two unit-3 examples (academy/src/finals/), from the same seeded data the site
generates, so learners can open the end result: real Excel workbooks (the stacked data with VLOOKUP formulas, pivot
tables, a report with formulas, charts, a dashboard sheet) and, for example 2, a PowerPoint deck and its PDF.

Needs LibreOffice (soffice, with its Python UNO bridge) and python-pptx. Run from the repository root:
    python3 academy/tools/make-finals.py
then rebuild the site (node academy/build.mjs), which embeds academy/src/finals/* in AI_Academy.html.
Every figure comes from the generators; data.test.mjs checks the finished files against them."""
import json, os, subprocess, sys, tempfile, time
import uno
from com.sun.star.beans import PropertyValue
from com.sun.star.table import CellAddress
from com.sun.star.awt import Rectangle
from com.sun.star.sheet.DataPilotFieldOrientation import ROW, COLUMN, DATA
from com.sun.star.sheet.GeneralFunction import SUM
from com.sun.star.chart.ChartDataRowSource import COLUMNS

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "academy", "src", "finals")
INK, ACC, SOFT, MINT, WHITE, LINE = 0x1F2937, 0x2F5FE6, 0xE8EEFD, 0x1F9D6B, 0xFFFFFF, 0xD7DCE5


def pv(name, value):
    p = PropertyValue(); p.Name = name; p.Value = value; return p


def col(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26); s = chr(65 + r) + s
    return s


class Office:
    def __enter__(self):
        self.proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--nologo", "--norestore", "--accept=pipe,name=academyfinals;urp;"], stderr=subprocess.DEVNULL)
        local = uno.getComponentContext()
        res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(120):
            try:
                self.ctx = res.resolve("uno:pipe,name=academyfinals;urp;StarOffice.ComponentContext"); break
            except Exception:
                time.sleep(0.5)
        self.desk = self.ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", self.ctx)
        return self

    def new_calc(self):
        return self.desk.loadComponentFromURL("private:factory/scalc", "_blank", 0, (pv("Hidden", True),))

    def open(self, path):
        return self.desk.loadComponentFromURL(uno.systemPathToFileUrl(path), "_blank", 0, (pv("Hidden", True),))

    def __exit__(self, *a):
        try: self.desk.terminate()
        except Exception: pass
        self.proc.wait(timeout=60)


class Book:
    """A small helper over one Calc document: sheets, values, formulas, formats, pivots, charts."""
    def __init__(self, office, names):
        self.doc = office.new_calc(); sh = self.doc.Sheets
        sh.getByIndex(0).Name = names[0]
        for i, n in enumerate(names[1:], 1): sh.insertNewByName(n, i)
        self.fmts = {}

    def sheet(self, name): return self.doc.Sheets.getByName(name)

    def fmt(self, code):
        if code not in self.fmts:
            nf, loc = self.doc.NumberFormats, uno.createUnoStruct("com.sun.star.lang.Locale")
            loc.Language, loc.Country = "en", "US"
            k = nf.queryKey(code, loc, False)
            self.fmts[code] = k if k != -1 else nf.addNew(code, loc)
        return self.fmts[code]

    def put(self, sheet, top_left, rows):   # values: numbers and text; a string starting with = is a formula
        s = self.sheet(sheet); c0, r0 = top_left
        width = max(len(r) for r in rows)
        rng = s.getCellRangeByPosition(c0, r0, c0 + width - 1, r0 + len(rows) - 1)
        if any(isinstance(v, str) and v.startswith("=") for r in rows for v in r):
            rng.setFormulaArray(tuple(tuple("" if v is None else str(v) for v in list(r) + [None] * (width - len(r))) for r in rows))
        else:
            rng.setDataArray(tuple(tuple("" if v is None else v for v in list(r) + [None] * (width - len(r))) for r in rows))
        return rng

    def style(self, sheet, ref, bold=False, bg=None, color=None, size=None, fmt=None, align=None):
        r = self.sheet(sheet).getCellRangeByName(ref)
        if bold: r.CharWeight = 150
        if bg is not None: r.CellBackColor = bg
        if color is not None: r.CharColor = color
        if size: r.CharHeight = size
        if fmt: r.NumberFormat = self.fmt(fmt)
        if align: r.HoriJustify = {"left": 1, "center": 2, "right": 3}[align]

    def widths(self, sheet, cm):
        cols = self.sheet(sheet).Columns
        for i, w in enumerate(cm): cols.getByIndex(i).Width = int(w * 1000)

    def head(self, sheet, ref):
        self.style(sheet, ref, bold=True, bg=ACC, color=WHITE)

    def pivot(self, sheet, name, at, source_sheet, source_ref, rows=(), cols=(), data=()):
        s = self.sheet(sheet); tables = s.getDataPilotTables(); d = tables.createDataPilotDescriptor()
        src = self.sheet(source_sheet).getCellRangeByName(source_ref)
        d.setSourceRange(src.getRangeAddress())
        try: d.ShowFilterButton = False
        except Exception: pass
        heads = [src.getCellByPosition(i, 0).getString() for i in range(src.getColumns().getCount())]
        fields = d.getDataPilotFields()
        for n in rows: fields.getByIndex(heads.index(n)).Orientation = ROW
        for n in cols: fields.getByIndex(heads.index(n)).Orientation = COLUMN
        for n in data:
            f = fields.getByIndex(heads.index(n)); f.Orientation = DATA; f.Function = SUM
        a = CellAddress(); a.Sheet = list(self.doc.Sheets.ElementNames).index(sheet); a.Column, a.Row = at
        tables.insertNewByName(name, a, d)

    def chart(self, sheet, name, ref, rect_cm, kind="column", title="", stacked=False, legend=False):
        s = self.sheet(sheet); r = Rectangle()
        r.X, r.Y, r.Width, r.Height = [int(v * 1000) for v in rect_cm]
        s.Charts.addNewByName(name, r, tuple(self.sheet_range(x).getRangeAddress() for x in ref.split(",")), True, True)
        ch = s.Charts.getByName(name).getEmbeddedObject()
        if kind == "bar":
            dg = ch.createInstance("com.sun.star.chart.BarDiagram"); ch.setDiagram(dg); dg.Vertical = True
        elif kind == "line":
            ch.setDiagram(ch.createInstance("com.sun.star.chart.LineDiagram"))
        dg = ch.getDiagram(); dg.DataRowSource = COLUMNS
        try:   # values from zero, on a white wall
            dg.YAxis.AutoMin = False; dg.YAxis.Min = 0; dg.Wall.FillColor = WHITE
        except Exception: pass
        if stacked: dg.Stacked = True
        if title:
            ch.HasMainTitle = True; ch.getTitle().String = title
        ch.HasLegend = legend

    def sheet_range(self, ref):
        sh, cells = ref.split("!")
        return self.sheet(sh).getCellRangeByName(cells)

    def save(self, path):
        self.doc.calculateAll()
        for s in self.doc.Sheets:
            for t in s.getDataPilotTables(): t.refresh()
        self.doc.calculateAll()
        self.doc.storeToURL(uno.systemPathToFileUrl(path), (pv("FilterName", "Calc MS Excel 2007 XML"),))
        self.doc.close(True)


def example1(office, D):
    E = D["ex1"]["expected"]; serial0 = D["serial0"]
    b = Book(office, ["Dashboard", "Report", "Pivots", "Data", "Models", "Countries"])
    b.put("Models", (0, 0), [["ModelCode", "ModelName"]] + D["models"]); b.head("Models", "A1:B1"); b.widths("Models", [3.2, 5.2])
    b.put("Countries", (0, 0), [["CountryCode", "CountryName"]] + D["countries"]); b.head("Countries", "A1:B1"); b.widths("Countries", [3.2, 5.2])
    rows = []
    for w in D["ex1"]["weeks"]:
        for r in w["rows"]:
            d, wo, line, mc, cc, q, v = r
            month = "2026-%02d" % ((__import__("datetime").date(1899, 12, 30) + __import__("datetime").timedelta(days=d)).month)
            rows.append([d, wo, line, mc, None, cc, None, q, v, month, w["file"]])
    n = len(rows); last = n + 1
    head = ["Date", "WorkOrder", "LineCode", "ModelCode", "ModelName", "CountryCode", "CountryName", "Qty", "Value", "Month", "SourceFile"]
    b.put("Data", (0, 0), [head])
    b.put("Data", (0, 1), [[r[0], r[1], r[2], r[3]] for r in rows])
    b.put("Data", (5, 1), [[r[5]] for r in rows])
    b.put("Data", (7, 1), [[r[7], r[8], r[9], r[10]] for r in rows])
    b.put("Data", (4, 1), [["=VLOOKUP(D%d;$Models.$A$2:$B$13;2;0)" % (i + 2)] for i in range(n)])
    b.put("Data", (6, 1), [["=VLOOKUP(F%d;$Countries.$A$2:$B$11;2;0)" % (i + 2)] for i in range(n)])
    b.head("Data", "A1:K1"); b.style("Data", "A2:A%d" % last, fmt="YYYY-MM-DD"); b.style("Data", "H2:I%d" % last, fmt="#,##0")
    b.style("Data", "E2:E%d" % last, bg=SOFT); b.style("Data", "G2:G%d" % last, bg=SOFT)
    b.widths("Data", [2.6, 3.0, 2.2, 2.6, 4.6, 2.8, 4.6, 1.8, 3.0, 2.2, 3.6])
    rng = lambda c: "$Data.$%s$2:$%s$%d" % (c, c, last)
    # the report: every figure is a formula over the Data sheet
    R = "Report"
    b.put(R, (0, 0), [["Production report - third quarter 2026 (1 July to 30 September)"]]); b.style(R, "A1", bold=True, size=16, color=INK)
    b.put(R, (0, 2), [["Total sets", "Total value (EGP)", "Export share", "Top model"],
                      ["=SUM(%s)" % rng("H"), "=SUM(%s)" % rng("I"), "=1-SUMIF(%s;\"Egypt\";%s)/SUM(%s)" % (rng("G"), rng("H"), rng("H")), "=INDEX(A26:A37;MATCH(MAX(B26:B37);B26:B37;0))"]])
    b.head(R, "A3:D3"); b.style(R, "A4:D4", bold=True, size=14, color=ACC); b.style(R, "A4:B4", fmt="#,##0"); b.style(R, "C4", fmt="0.0%")
    months = ["2026-07", "2026-08", "2026-09"]
    b.put(R, (0, 5), [["Month", "Sets", "Value (EGP)"]] + [[m, "=SUMIF(%s;A%d;%s)" % (rng("J"), 7 + i, rng("H")), "=SUMIF(%s;A%d;%s)" % (rng("J"), 7 + i, rng("I"))] for i, m in enumerate(months)] + [["Total", "=SUM(B7:B9)", "=SUM(C7:C9)"]])
    lines = sorted(E["lines"])
    b.put(R, (4, 5), [["Line", "Sets"]] + [[l, "=SUMIF(%s;E%d;%s)" % (rng("C"), 7 + i, rng("H"))] for i, l in enumerate(lines)] + [["Total", "=SUM(F7:F12)"]])
    countries = [c for c, _ in sorted(E["countries"].items(), key=lambda x: -x[1])]
    b.put(R, (0, 13), [["Country", "Sets", "Value (EGP)"]] + [[c, "=SUMIF(%s;A%d;%s)" % (rng("G"), 15 + i, rng("H")), "=SUMIF(%s;A%d;%s)" % (rng("G"), 15 + i, rng("I"))] for i, c in enumerate(countries)])
    models = [m for m, _ in sorted(E["models"].items(), key=lambda x: -x[1])]
    b.put(R, (0, 24), [["Model", "Sets"]] + [[m, "=SUMIF(%s;A%d;%s)" % (rng("E"), 26 + i, rng("H"))] for i, m in enumerate(models)])
    for ref in ("A6:C6", "E6:F6", "A14:C14", "A25:B25"): b.head(R, ref)
    for ref in ("B7:C10", "F7:F13", "B15:C24", "B26:B37"): b.style(R, ref, fmt="#,##0")
    for ref in ("A10:C10", "E13:F13"): b.style(R, ref, bold=True, bg=SOFT)
    b.widths(R, [5.2, 3.6, 4.0, 3.4, 2.6, 3.0])
    # pivot tables over the Data sheet
    src = "A1:K%d" % last
    b.put("Pivots", (0, 0), [["Pivot tables over the Data sheet (Data > Refresh All updates them)"]]); b.style("Pivots", "A1", bold=True, size=13)
    b.pivot("Pivots", "PT_MonthLine", (0, 2), "Data", src, rows=["Month"], cols=["LineCode"], data=["Qty"])
    b.pivot("Pivots", "PT_Country", (0, 12), "Data", src, rows=["CountryName"], data=["Qty", "Value"])
    b.pivot("Pivots", "PT_Model", (5, 12), "Data", src, rows=["ModelName"], data=["Qty"])
    b.widths("Pivots", [5.0, 2.6, 3.2, 2.4, 2.4, 5.0, 2.6, 2.4])
    # the dashboard: four figures and four charts, all read from the pivot tables (GETPIVOTDATA), so slicers connected
    # to the pivot tables filter every figure and every chart on it; the Report sheet stays the full-quarter printout
    Dd = "Dashboard"
    b.put(Dd, (0, 0), [["Production dashboard - Q3 2026"]]); b.style(Dd, "A1", bold=True, size=18, color=INK)
    b.put(Dd, (0, 1), [["Slicers: click a pivot table on the Pivots sheet, then PivotTable Analyze > Insert Slicer (Month, LineCode, CountryName), then Report Connections: tick all three tables."]])
    b.style(Dd, "A2", color=0x6B7280)
    PM, PC, PD = "$Pivots.$A$3", "$Pivots.$A$13", "$Pivots.$F$13"
    gp = lambda field, pt, k=None, item=None: ('IFERROR(GETPIVOTDATA("%s";%s;"%s";%s);0)' % (field, pt, k, item)) if k else 'GETPIVOTDATA("%s";%s)' % (field, pt)
    b.put(Dd, (9, 2), [["Chart data, read from the pivot tables (the slicers filter it)"]]); b.style(Dd, "J3", bold=True, color=0x6B7280)
    b.put(Dd, (9, 4), [["Month", "Sets"]] + [[m, "=" + gp("Qty", PM, "Month", "J%d" % (6 + i))] for i, m in enumerate(months)])
    b.put(Dd, (12, 4), [["Line", "Sets"]] + [[l, "=" + gp("Qty", PM, "LineCode", "M%d" % (6 + i))] for i, l in enumerate(lines)])
    b.put(Dd, (9, 13), [["Country", "Sets"]] + [[c, "=" + gp("Qty", PC, "CountryName", "J%d" % (15 + i))] for i, c in enumerate(countries)])
    b.put(Dd, (12, 13), [["Model", "Sets"]] + [[m, "=" + gp("Qty", PD, "ModelName", "M%d" % (15 + i))] for i, m in enumerate(models)])
    for ref in ("J5:K5", "M5:N5", "J14:K14", "M14:N14"): b.head(Dd, ref)
    for ref in ("K6:K8", "N6:N11", "K15:K24", "N15:N26"): b.style(Dd, ref, fmt="#,##0")
    b.put(Dd, (0, 2), [["Total sets", "", "Value (EGP million)", "", "Export share", "", "Top model"],
                       ["=" + gp("Qty", PC), "", "=" + gp("Value", PC) + "/1000000", "", "=1-" + gp("Qty", PC, "CountryName", '"Egypt"') + "/A4", "", "=INDEX($M$15:$M$26;MATCH(MAX($N$15:$N$26);$N$15:$N$26;0))"]])
    b.style(Dd, "A3:G3", bold=True, color=0x6B7280); b.style(Dd, "A4:G4", bold=True, size=16, color=ACC)
    b.style(Dd, "A4", fmt="#,##0"); b.style(Dd, "C4", fmt="#,##0.0"); b.style(Dd, "E4", fmt="0.0%")
    b.widths(Dd, [3.4, 0.6, 4.4, 0.6, 3.4, 0.6, 5.0, 0.6, 5.9, 4.6, 2.4, 0.6, 4.6, 2.4])
    b.chart(Dd, "SetsByMonth", "Dashboard!J5:K8", (0.3, 2.2, 11.5, 7.2), "column", "Sets by month")
    b.chart(Dd, "SetsByLine", "Dashboard!M5:N11", (12.2, 2.2, 11.5, 7.2), "column", "Sets by line")
    b.chart(Dd, "SetsByCountry", "Dashboard!J14:K24", (0.3, 9.8, 11.5, 8.0), "bar", "Sets by country")
    b.chart(Dd, "SetsByModel", "Dashboard!M14:N26", (12.2, 9.8, 11.5, 8.0), "bar", "Sets by model")
    b.save(os.path.join(OUT, "Production_Q3_Final.xlsx"))


def example2(office, D):
    import datetime
    X = D["ex2"]; E = X["expected"]
    b = Book(office, ["Dashboard", "Report", "Pivots", "Sales", "Shipments", "Returns", "Targets", "Models", "Countries"])
    b.put("Models", (0, 0), [["ModelCode", "ModelName"]] + D["models"]); b.head("Models", "A1:B1"); b.widths("Models", [3.2, 5.2])
    b.put("Countries", (0, 0), [["CountryCode", "CountryName"]] + D["countries"]); b.head("Countries", "A1:B1"); b.widths("Countries", [3.2, 5.2])
    month = lambda serial: "2026-%02d" % (datetime.date(1899, 12, 30) + datetime.timedelta(days=serial)).month
    sales = []
    for ch in ("Cairo", "Alexandria", "Delta"):
        for d, inv, mc, q, v in X["branches"][ch]: sales.append([month(d), ch, "EG", None, mc, None, q, v, "Sales_%s_Q3.xlsx" % ch])
    for m, cc, mc, q, v in X["exports"]: sales.append(["2026-%02d" % m, "Export", cc, None, mc, None, q, v, "Export_Sales_2026-%02d.pdf" % m])
    n = len(sales); last = n + 1
    b.put("Sales", (0, 0), [["Month", "Channel", "CountryCode", "CountryName", "ModelCode", "ModelName", "Qty", "NetSales", "SourceFile"]])
    b.put("Sales", (0, 1), [r[:3] for r in sales]); b.put("Sales", (4, 1), [[r[4]] for r in sales]); b.put("Sales", (6, 1), [r[6:] for r in sales])
    b.put("Sales", (3, 1), [["=VLOOKUP(C%d;$Countries.$A$2:$B$11;2;0)" % (i + 2)] for i in range(n)])
    b.put("Sales", (5, 1), [["=VLOOKUP(E%d;$Models.$A$2:$B$13;2;0)" % (i + 2)] for i in range(n)])
    b.head("Sales", "A1:I1"); b.style("Sales", "G2:H%d" % last, fmt="#,##0"); b.style("Sales", "D2:D%d" % last, bg=SOFT); b.style("Sales", "F2:F%d" % last, bg=SOFT)
    b.widths("Sales", [2.2, 2.8, 2.8, 4.6, 2.6, 4.6, 1.8, 3.2, 5.6])
    b.put("Shipments", (0, 0), [["Month", "Channel", "Shipments", "Units", "FreightEGP"]] + [["2026-%02d" % m, ch, k, u, f] for m, ch, k, u, f in X["ships"]])
    b.head("Shipments", "A1:E1"); b.style("Shipments", "C2:E13", fmt="#,##0"); b.widths("Shipments", [2.4, 3.0, 2.6, 2.4, 3.0])
    ret = [["2026-%02d" % m, mc, "=VLOOKUP(B%d;$Models.$A$2:$B$13;2;0)" % (i + 2), q] for i, (m, mc, q) in enumerate(X["returns"])]
    b.put("Returns", (0, 0), [["Month", "ModelCode", "ModelName", "Units"]] + ret); b.head("Returns", "A1:D1"); b.style("Returns", "C2:C%d" % (len(ret) + 1), bg=SOFT); b.widths("Returns", [2.4, 2.6, 4.6, 2.0])
    b.put("Targets", (0, 0), [["Channel", "TargetUnits", "TargetSales"]] + X["targets"]); b.head("Targets", "A1:C1"); b.style("Targets", "B2:C5", fmt="#,##0"); b.widths("Targets", [3.0, 3.0, 3.6])
    S = lambda c: "$Sales.$%s$2:$%s$%d" % (c, c, last)
    R = "Report"; chs = ["Cairo", "Alexandria", "Delta", "Export"]
    b.put(R, (0, 0), [["Commercial review - third quarter 2026"]]); b.style(R, "A1", bold=True, size=16, color=INK)
    b.put(R, (0, 2), [["Channel", "Sets", "Net sales (EGP)", "Target (EGP)", "Achievement", "Units shipped", "Freight (EGP)", "Freight per set"]] +
          [[c, "=SUMIF(%s;A%d;%s)" % (S("B"), 4 + i, S("G")), "=SUMIF(%s;A%d;%s)" % (S("B"), 4 + i, S("H")), "=VLOOKUP(A%d;$Targets.$A$2:$C$5;3;0)" % (4 + i), "=C%d/D%d" % (4 + i, 4 + i),
            "=SUMIF($Shipments.$B$2:$B$13;A%d;$Shipments.$D$2:$D$13)" % (4 + i), "=SUMIF($Shipments.$B$2:$B$13;A%d;$Shipments.$E$2:$E$13)" % (4 + i), "=G%d/F%d" % (4 + i, 4 + i)] for i, c in enumerate(chs)] +
          [["Total", "=SUM(B4:B7)", "=SUM(C4:C7)", "=SUM(D4:D7)", "=C8/D8", "=SUM(F4:F7)", "=SUM(G4:G7)", "=G8/F8"]])
    b.head(R, "A3:H3"); b.style(R, "B4:D8", fmt="#,##0"); b.style(R, "E4:E8", fmt="0.0%"); b.style(R, "F4:G8", fmt="#,##0"); b.style(R, "H4:H8", fmt="#,##0.0"); b.style(R, "A8:H8", bold=True, bg=SOFT)
    b.put(R, (0, 10), [["Month", "Net sales (EGP)"]] + [[m, "=SUMIF(%s;A%d;%s)" % (S("A"), 12 + i, S("H"))] for i, m in enumerate(["2026-07", "2026-08", "2026-09"])])
    b.head(R, "A11:B11"); b.style(R, "B12:B14", fmt="#,##0")
    countries = [c for c, _ in sorted(E["countries"].items(), key=lambda x: -x[1])]
    b.put(R, (3, 10), [["Export country", "Net sales (EGP)"]] + [[c, "=SUMIFS(%s;%s;\"Export\";%s;D%d)" % (S("H"), S("B"), S("D"), 12 + i)] for i, c in enumerate(countries)])
    b.head(R, "D11:E11"); b.style(R, "E12:E20", fmt="#,##0")
    rr = E["returnRates"]
    b.put(R, (0, 22), [["Model", "ModelName", "Returned", "Sold", "Return rate"]] + [[x["mc"], "=VLOOKUP(A%d;$Models.$A$2:$B$13;2;0)" % (24 + i), "=SUMIF($Returns.$B$2:$B$37;A%d;$Returns.$D$2:$D$37)" % (24 + i), "=SUMIF(%s;A%d;%s)" % (S("E"), 24 + i, S("G")), "=C%d/D%d" % (24 + i, 24 + i)] for i, x in enumerate(rr)])
    b.head(R, "A23:E23"); b.style(R, "C24:D35", fmt="#,##0"); b.style(R, "E24:E35", fmt="0.0%"); b.style(R, "A24:E26", bold=True)
    b.put(R, (9, 2), [["Channel", "Net sales (EGP M)", "Target (EGP M)", "Freight per set"]] + [["=A%d" % (4 + i), "=C%d/1000000" % (4 + i), "=D%d/1000000" % (4 + i), "=H%d" % (4 + i)] for i in range(4)])
    b.head(R, "J3:M3"); b.style(R, "K4:M7", fmt="#,##0.0"); b.put(R, (9, 1), [["Chart data"]]); b.style(R, "J2", bold=True, color=0x6B7280)
    b.widths(R, [3.6, 4.4, 4.0, 3.6, 3.2, 3.2, 3.2, 3.0, 0.6, 3.0, 3.6, 3.4, 3.2])
    src = "A1:I%d" % last
    b.put("Pivots", (0, 0), [["Pivot tables over the Sales sheet (Data > Refresh All updates them)"]]); b.style("Pivots", "A1", bold=True, size=13)
    b.pivot("Pivots", "PT_ChannelMonth", (0, 2), "Sales", src, rows=["Channel"], cols=["Month"], data=["NetSales"])
    b.pivot("Pivots", "PT_Country", (0, 11), "Sales", src, rows=["CountryName"], data=["Qty", "NetSales"])
    b.pivot("Pivots", "PT_Model", (5, 11), "Sales", src, rows=["ModelName"], data=["Qty", "NetSales"])
    b.widths("Pivots", [5.0, 3.2, 3.2, 3.2, 3.2, 5.0, 2.6, 3.4])
    Dd = "Dashboard"
    b.put(Dd, (0, 0), [["Commercial review - Q3 2026"]]); b.style(Dd, "A1", bold=True, size=18, color=INK)
    b.put(Dd, (0, 2), [["Net sales (EGP million)", "", "Sets sold", "", "Achievement", "", "Freight per set (EGP)"],
                       ["=Report.C8/1000000", "", "=Report.B8", "", "=Report.E8", "", "=Report.H8"]])
    b.style(Dd, "A3:G3", bold=True, color=0x6B7280); b.style(Dd, "A4:G4", bold=True, size=16, color=ACC)
    b.style(Dd, "A4", fmt="#,##0.0"); b.style(Dd, "C4", fmt="#,##0"); b.style(Dd, "E4", fmt="0.0%"); b.style(Dd, "G4", fmt="#,##0.0")
    b.widths(Dd, [4.6, 0.6, 3.4, 0.6, 3.4, 0.6, 4.6])
    b.chart(Dd, "SalesVsTarget", "Report!J3:L7", (0.3, 2.2, 11.5, 7.2), "column", "Net sales and target by channel (EGP M)", legend=True)
    b.chart(Dd, "Trend", "Report!A11:B14", (12.2, 2.2, 11.5, 7.2), "column", "Net sales by month")
    b.chart(Dd, "ExportByCountry", "Report!D11:E20", (0.3, 9.8, 11.5, 8.0), "bar", "Export sales by country")
    b.chart(Dd, "FreightPerSet", "Report!J3:J7,Report!M3:M7", (12.2, 9.8, 11.5, 8.0), "column", "Freight per set by channel (EGP)")
    b.save(os.path.join(OUT, "Q3_Commercial_Final.xlsx"))


def deck(D, path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    E = D["ex2"]["expected"]; chs = ["Cairo", "Alexandria", "Delta", "Export"]
    acc, ink, grey = RGBColor(0x2F, 0x5F, 0xE6), RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80)
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, sub=None):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = acc; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9)).text_frame
        tb.text = title; p = tb.paragraphs[0]; p.runs[0].font.size = Pt(30); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = ink
        if sub:
            q = tb.add_paragraph(); q.text = sub; q.runs[0].font.size = Pt(14); q.runs[0].font.color.rgb = grey
        return s

    def chart(s, kind, cats, series, box, fmt="#,##0", legend=False):
        cd = CategoryChartData(); cd.categories = cats
        for name, vals in series: cd.add_series(name, vals)
        g = s.shapes.add_chart(kind, *box, cd).chart
        g.has_legend = legend
        if legend: g.legend.position = XL_LEGEND_POSITION.BOTTOM; g.legend.include_in_layout = False
        pl = g.plots[0]; pl.has_data_labels = True; pl.data_labels.number_format = fmt; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(11)
        for i, sr in enumerate(pl.series):
            sr.format.fill.solid(); sr.format.fill.fore_color.rgb = acc if i == 0 else RGBColor(0xB8, 0xC7, 0xF5)
        return g

    def figures(s, items, top=Inches(1.7)):
        w = Inches(2.9)
        for i, (big, small) in enumerate(items):
            box = s.shapes.add_shape(1, Inches(0.6) + i * Inches(3.1), top, w, Inches(1.5)); box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xE8, 0xEE, 0xFD); box.line.fill.background()
            tf = box.text_frame; tf.text = big; tf.paragraphs[0].runs[0].font.size = Pt(28); tf.paragraphs[0].runs[0].font.bold = True; tf.paragraphs[0].runs[0].font.color.rgb = acc
            p = tf.add_paragraph(); p.text = small; p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = ink

    m = lambda v: "{:,.1f}".format(v / 1e6)
    slide("Q3 Commercial Review", "1 July to 30 September 2026 · prepared from the branch files, the system reports and the memos")
    s = slide("Key figures")
    figures(s, [("EGP %s M" % m(E["salesValue"]), "Net sales"), ("{:,}".format(E["salesQty"]), "Sets sold"),
                ("%.1f%%" % (E["salesValue"] / sum(E["target"][c]["v"] for c in chs) * 100), "Achievement against target"), ("EGP %.1f" % (E["freight"] / E["shipUnits"]), "Freight per set")])
    s = slide("Sales against target by channel", "Net sales and target, EGP million")
    chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, chs, [("Net sales", [round(E["channels"][c]["v"] / 1e6, 1) for c in chs]), ("Target", [round(E["target"][c]["v"] / 1e6, 1) for c in chs])], (Inches(0.6), Inches(1.6), Inches(8.2), Inches(5.4)), "#,##0.0", True)
    tb = s.shapes.add_textbox(Inches(9.2), Inches(2.0), Inches(3.6), Inches(4)).text_frame
    for i, c in enumerate(chs):
        p = tb.paragraphs[0] if i == 0 else tb.add_paragraph(); p.text = "%s: %.1f%%" % (c, E["achievement"][c]); p.runs[0].font.size = Pt(22); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = ink; p.space_after = Pt(14)
    s = slide("Monthly trend", "Net sales, EGP million")
    chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, ["July", "August", "September"], [("Net sales", [round(E["months"][k] / 1e6, 1) for k in ("7", "8", "9")])], (Inches(0.6), Inches(1.6), Inches(12), Inches(5.4)), "#,##0.0")
    s = slide("Export sales by country", "Net sales, EGP million")
    cs = sorted(E["countries"].items(), key=lambda x: x[1])
    chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [c for c, _ in cs], [("Net sales", [round(v / 1e6, 1) for _, v in cs])], (Inches(0.6), Inches(1.6), Inches(12), Inches(5.4)), "#,##0.0")
    s = slide("Logistics", "%d shipments · %s sets shipped · freight per set by channel, EGP" % (E["shipments"], "{:,}".format(E["shipUnits"])))
    chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, chs, [("Freight per set", [E["freightPerUnit"][c] for c in chs])], (Inches(0.6), Inches(1.6), Inches(12), Inches(5.4)), "#,##0.0")
    s = slide("Returns", "The three models with the highest return rate")
    figures(s, [("%.1f%%" % x["rate"], "%s · %s" % (x["mc"], x["name"])) for x in E["topReturns"]])
    tb = s.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(12), Inches(1)).text_frame
    tb.text = "%s sets returned in the quarter, against %s sold." % ("{:,}".format(E["returns"]), "{:,}".format(E["salesQty"])); tb.paragraphs[0].runs[0].font.size = Pt(16); tb.paragraphs[0].runs[0].font.color.rgb = grey
    slide("For discussion", "The director's conclusions and decisions")
    prs.save(path)


def main():
    data = json.loads(subprocess.check_output(["node", os.path.join(ROOT, "academy", "tools", "export-data.mjs")]))
    os.makedirs(OUT, exist_ok=True)
    with Office() as office:
        example1(office, data)
        example2(office, data)
    pptx = os.path.join(OUT, "Q3_Commercial_Review.pptx")
    deck(data, pptx)
    with tempfile.TemporaryDirectory() as tmp:   # the deck as PDF, through LibreOffice Impress
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        os.replace(os.path.join(tmp, "Q3_Commercial_Review.pdf"), os.path.join(OUT, "Q3_Commercial_Review.pdf"))
    for f in sorted(os.listdir(OUT)): print("%-32s %8d bytes" % (f, os.path.getsize(os.path.join(OUT, f))))


if __name__ == "__main__":
    main()
